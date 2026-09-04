"""Generate Pengujian_BlackBox.docx — dokumen pengujian black box untuk aplikasi Skill Gap Detection.

Dapat dijalankan ulang untuk menghasilkan dokumen baru (misal setelah pengujian selesai).
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Pengujian_BlackBox.docx")

HEADER_FILL = "1F4E5F"
ALT_FILL = "EAF2F5"

COLUMNS = ["ID", "Skenario Pengujian", "Input", "Langkah Uji", "Hasil yang Diharapkan", "Hasil Aktual", "Status", "Catatan"]
COL_WIDTHS = [Cm(1.6), Cm(3.4), Cm(3.4), Cm(4.6), Cm(5.2), Cm(3.4), Cm(1.9), Cm(2.2)]

API_CASES = [
    (
        "TC-API-01",
        "Memeriksa kesehatan API (GET /health)",
        "-",
        "1. Jalankan backend\n2. Buka http://localhost:8000/health",
        "HTTP 200. JSON berisi status=\"ok\", model_loaded=true, categories_available > 0",
    ),
    (
        "TC-API-02",
        "Melihat daftar kategori pekerjaan (GET /job-categories)",
        "-",
        "1. Buka http://localhost:8000/job-categories",
        "HTTP 200. JSON berisi count=24 dan daftar kategori pekerjaan IT",
    ),
    (
        "TC-API-03",
        "Deteksi gap dengan data valid (POST /detect-gap)",
        '{"skills": ["Python", "SQL", "Data Analysis"], "target_job": "Data Science & AI", "top_n": 15}',
        "1. Buka http://localhost:8000/docs\n2. Eksekusi endpoint /detect-gap dengan input di samping",
        "HTTP 200. Response berisi target_job, benchmark_count=15, matched, gap, gap_score, match_score (0-100)",
    ),
    (
        "TC-API-04",
        "Deteksi gap dengan target_job tidak dikenal",
        '{"skills": ["Python"], "target_job": "Astronaut", "top_n": 15}',
        "1. Eksekusi /detect-gap dengan target_job yang tidak ada di daftar",
        "HTTP 404. Pesan detail berisi daftar kategori yang tersedia",
    ),
    (
        "TC-API-05",
        "Deteksi gap dengan huruf campuran besar/kecil",
        '{"skills": ["Python"], "target_job": "data science & AI", "top_n": 15}',
        "1. Eksekusi /detect-gap dengan target_job yang hurufnya tidak sama persis dengan katalog",
        "HTTP 200. target_job dinormalisasi menjadi \"Data Science & AI\" (case-insensitive)",
    ),
    (
        "TC-API-06",
        "Deteksi gap dengan daftar skill kosong",
        '{"skills": [], "target_job": "Software Development", "top_n": 15}',
        "1. Eksekusi /detect-gap dengan array skills kosong",
        "HTTP 200. Semua skill masuk ke daftar gap, gap_score=100, match_score=0",
    ),
    (
        "TC-API-07",
        "Deteksi gap dengan top_n sangat besar",
        '{"skills": ["Python"], "target_job": "Data Engineering", "top_n": 300}',
        "1. Eksekusi /detect-gap dengan top_n=300",
        "HTTP 200. benchmark_count tidak melebihi jumlah total kandidat skill model, tidak terjadi error",
    ),
    (
        "TC-API-08",
        "Deteksi gap dengan top_n = 0",
        '{"skills": ["Python"], "target_job": "Data Engineering", "top_n": 0}',
        "1. Eksekusi /detect-gap dengan top_n=0",
        "HTTP 200. benchmark_count=0, gap dan matched kosong, gap_score=0",
    ),
    (
        "TC-API-09",
        "Mendapatkan rekomendasi skill valid (POST /recommend)",
        '{"skills": ["Python", "SQL"], "target_job": "Data Analytics & BI", "top_n": 10}',
        "1. Eksekusi endpoint /recommend dengan input di samping",
        "HTTP 200. recommendations berisi 10 item dengan priority_rank berurutan 1..10, tiap item punya skill dan prob",
    ),
    (
        "TC-API-10",
        "Rekomendasi dengan target_job tidak dikenal",
        '{"skills": ["Python"], "target_job": "Astronaut", "top_n": 10}',
        "1. Eksekusi /recommend dengan target_job tidak valid",
        "HTTP 404. Pesan detail berisi daftar kategori yang tersedia",
    ),
    (
        "TC-API-11",
        "Parsing CV PDF valid (POST /parse-cv)",
        "File PDF berisi teks skill (mis. Python, SQL, Machine Learning)",
        "1. Upload file PDF valid dengan multipart/form-data field \"file\"",
        "HTTP 200. detected_skills berisi skill yang dikenali, skill_count sesuai, extraction_notes dan raw_text_preview ada",
    ),
    (
        "TC-API-12",
        "Parsing file non-PDF (mis. .txt)",
        "File dengan ekstensi .txt atau .docx",
        "1. Upload file non-PDF",
        "HTTP 400. Pesan \"Only PDF files are supported. Please upload a .pdf file.\"",
    ),
    (
        "TC-API-13",
        "Parsing PDF berukuran > 10 MB",
        "File PDF > 10 MB",
        "1. Upload file PDF berukuran lebih dari 10 MB",
        "HTTP 413. Pesan \"File too large. Maximum size is 10 MB.\"",
    ),
    (
        "TC-API-14",
        "Parsing PDF korup/rusak",
        "File ber-ekstensi .pdf tetapi isinya korup (bukan PDF valid)",
        "1. Upload file PDF korup",
        "HTTP 422 (ValueError) atau HTTP 503 (RuntimeError) dengan pesan kesalahan sesuai",
    ),
    (
        "TC-API-15",
        "Parsing PDF tanpa teks (hasil scan gambar)",
        "File PDF berisi gambar/scan tanpa lapisan teks",
        "1. Upload file PDF hasil scan yang tidak mengandung teks",
        "HTTP 200. skill_count=0, detected_skills kosong, extraction_notes menyebutkan tidak ada teks terekstrak",
    ),
    (
        "TC-API-16",
        "Request /parse-cv tanpa file",
        "-",
        "1. Kirim request POST /parse-cv tanpa field file",
        "HTTP 422 (FastAPI validation error)",
    ),
]

UI_CASES = [
    (
        "TC-UI-01",
        "Membuka halaman utama aplikasi",
        "Akses http://localhost:5173",
        "1. Jalankan backend dan frontend\n2. Buka URL di browser",
        "Halaman Landing tampil dengan benar, judul dan menu navigasi terlihat",
    ),
    (
        "TC-UI-02",
        "Navigasi antar halaman",
        "-",
        "1. Klik setiap menu: Beranda, Scan CV, Manual Input, Self Assessment, Tentang",
        "Semua halaman dapat diakses dan menampilkan konten sesuai rute (/ , /scan-cv, /manual-input, /self-assessment, /about)",
    ),
    (
        "TC-UI-03",
        "Menambah skill melalui tombol Tambah",
        "Skill: \"Python\"",
        "1. Buka halaman Manual Input\n2. Ketik \"Python\" pada input\n3. Klik tombol Tambah",
        "Skill muncul sebagai chip dengan nomor urut 1, counter menampilkan \"1 skill ditambahkan\"",
    ),
    (
        "TC-UI-04",
        "Menambah skill melalui tombol Enter",
        "Skill: \"React\"",
        "1. Ketik \"React\" pada input\n2. Tekan Enter",
        "Skill ditambahkan tanpa perlu klik tombol, input dikosongkan kembali",
    ),
    (
        "TC-UI-05",
        "Mencegah penambahan skill duplikat",
        "Skill: \"Python\" ditambahkan dua kali",
        "1. Tambahkan \"Python\"\n2. Tambahkan \"Python\" lagi",
        "Skill duplikat tidak ditambahkan, hanya 1 chip \"Python\" yang tampil",
    ),
    (
        "TC-UI-06",
        "Menghapus skill yang sudah ditambahkan",
        "-",
        "1. Tambahkan 2 skill\n2. Klik tombol × pada salah satu chip",
        "Chip skill terhapus, counter dan penomoran chip menyesuaikan",
    ),
    (
        "TC-UI-07",
        "Tombol Tambah nonaktif saat input kosong",
        "-",
        "1. Buka halaman Manual Input tanpa mengetik apa pun",
        "Tombol Tambah berstatus disabled (tidak dapat diklik)",
    ),
    (
        "TC-UI-08",
        "Tombol analisis nonaktif tanpa skill/job",
        "-",
        "1. Buka halaman Manual Input tanpa menambahkan skill\n2. Periksa tombol \"Jalankan Analisis Kesenjangan\"",
        "Bagian pemilihan job dan tombol analisis tidak muncul sebelum ada skill, atau tombol disabled",
    ),
    (
        "TC-UI-09",
        "Analisis gap valid dari input manual",
        "Skills: Python, SQL, React — Target: Software Development",
        "1. Tambahkan skill\n2. Pilih target job\n3. Klik \"Jalankan Analisis Kesenjangan\"",
        "Hasil menampilkan gap score, match score, daftar matched, gap, dan rekomendasi prioritas belajar",
    ),
    (
        "TC-UI-10",
        "Menampilkan pesan error saat API gagal",
        "-",
        "1. Matikan backend\n2. Jalankan analisis dari halaman Manual Input",
        "Panel error merah tampil dengan pesan kesalahan (mis. network error)",
    ),
    (
        "TC-UI-11",
        "Upload CV PDF pada halaman Scan CV",
        "File CV berformat PDF",
        "1. Buka /scan-cv\n2. Pilih file PDF pada zona upload",
        "File terpilih ditampilkan (nama, ukuran), log \"FILE LOADED\" muncul, tombol \"Analisis CV Saya\" aktif",
    ),
    (
        "TC-UI-12",
        "Mengganti file CV yang sudah dipilih",
        "Dua file PDF berbeda",
        "1. Upload file pertama\n2. Klik \"← Ganti file\"\n3. Pilih file kedua",
        "File pertama diganti file kedua, log dan status di-reset",
    ),
    (
        "TC-UI-13",
        "Menjalankan analisis CV (flow 5 langkah)",
        "File CV PDF berisi skill",
        "1. Upload PDF\n2. Klik \"Analisis CV Saya\"\n3. Tunggu proses selesai",
        "Log proses ditampilkan berurutan, skill terdeteksi muncul sebagai tag, flow indicator maju ke langkah 3",
    ),
    (
        "TC-UI-14",
        "Mengedit skill hasil deteksi CV",
        "-",
        "1. Setelah scan, tambahkan 1 skill baru dan hapus 1 tag skill",
        "Tag skill dapat ditambah dan dihapus, daftar skill terupdate sebelum deteksi gap",
    ),
    (
        "TC-UI-15",
        "Deteksi gap dari hasil scan CV",
        "Target: Data Science & AI",
        "1. Upload + scan CV\n2. Pilih target job\n3. Klik \"Deteksi Kesenjangan\"",
        "Flow sampai langkah 5, hasil gap dan rekomendasi ditampilkan, tombol \"Ganti Target Job\" dan \"Upload CV Baru\" tersedia",
    ),
    (
        "TC-UI-16",
        "Upload file non-PDF pada Scan CV",
        "File ber-ekstensi .txt",
        "1. Coba pilih file .txt pada zona upload",
        "File ditolak atau saat dianalisis muncul pesan error \"Only PDF files are supported...\"",
    ),
    (
        "TC-UI-17",
        "Mengisi formulir self assessment",
        "-",
        "1. Buka /self-assessment\n2. Pilih level untuk tiap skill (0-3)\n3. Klik submit",
        "Form tervalidasi, setelah submit muncul bagian pemilihan target pekerjaan (Tahap 2)",
    ),
    (
        "TC-UI-18",
        "Analisis gap dari hasil self assessment",
        "Target: Software Development",
        "1. Isi form assessment\n2. Pilih target job\n3. Klik \"Jalankan Analisis Kesenjangan\"",
        "Hasil gap dan rekomendasi tampil, skill dengan level 0 tidak ikut dikirim sebagai skill dimiliki",
    ),
    (
        "TC-UI-19",
        "Membuka halaman Tentang",
        "-",
        "1. Buka /about atau klik menu Tentang",
        "Halaman berisi informasi tentang aplikasi tampil tanpa error",
    ),
]

INT_CASES = [
    (
        "TC-INT-01",
        "Kesesuaian skill hasil scan CV dengan daftar matched",
        "CV PDF berisi skill Python, SQL, Pandas — Target: Data Science & AI",
        "1. Scan CV pada /scan-cv, catat detected_skills\n2. Pilih target job\n3. Jalankan deteksi gap\n4. Bandingkan detected_skills dengan daftar matched",
        "Setiap skill hasil scan yang merupakan skill benchmark job muncul pada daftar matched (bukan gap)",
    ),
    (
        "TC-INT-02",
        "Kesesuaian daftar gap dengan daftar rekomendasi",
        "Input yang sama untuk /detect-gap dan /recommend",
        "1. Jalankan /detect-gap, catat daftar gap\n2. Jalankan /recommend dengan input yang sama\n3. Bandingkan kedua daftar",
        "Daftar gap pada /detect-gap identik (skill & urutan) dengan daftar recommendations pada /recommend",
    ),
    (
        "TC-INT-03",
        "Konsistensi jumlah matched + gap = benchmark_count",
        "Hasil respons /detect-gap",
        "1. Jalankan /detect-gap dengan data valid\n2. Hitung len(matched) + len(gap)\n3. Bandingkan dengan benchmark_count",
        "len(matched) + len(gap) = benchmark_count",
    ),
    (
        "TC-INT-04",
        "Konsistensi gap_score + match_score = 100",
        "Hasil respons /detect-gap",
        "1. Jalankan /detect-gap dengan data valid\n2. Jumlahkan gap_score dan match_score",
        "gap_score + match_score = 100 (dalam toleransi pembulatan 0,1)",
    ),
    (
        "TC-INT-05",
        "Urutan rekomendasi sesuai probabilitas tertinggi",
        "Hasil respons /recommend",
        "1. Jalankan /recommend dengan data valid\n2. Bandingkan nilai prob tiap item berurutan\n3. Periksa priority_rank",
        "priority_rank berurutan 1..n dan prob item selalu menurun (item rank 1 memiliki prob tertinggi)",
    ),
    (
        "TC-INT-06",
        "Determinisme hasil scan CV",
        "File CV PDF yang sama",
        "1. Scan CV yang sama dua kali berturut-turut\n2. Bandingkan detected_skills kedua hasil",
        "Hasil deteksi identik pada kedua percobaan (tidak ada perbedaan acak)",
    ),
    (
        "TC-INT-07",
        "Skill non-benchmark tidak menyebabkan error",
        "CV berisi skill di luar kandidat model (mis. \"Photoshop\")",
        "1. Scan CV berisi skill non-benchmark\n2. Jalankan deteksi gap\n3. Periksa respons",
        "Proses berhasil tanpa error; skill non-benchmark tidak muncul di matched maupun gap",
    ),
]

ASS_CASES = [
    (
        "TC-ASS-01",
        "Kelengkapan form penilaian",
        "Halaman /self-assessment",
        "1. Buka halaman Self Assessment\n2. Periksa subdomain dan daftar skill yang ditampilkan",
        "Form menampilkan 4 subdomain (Software Development, Data Engineering, Cybersecurity, DevOps & Cloud) dengan daftar skill masing-masing",
    ),
    (
        "TC-ASS-02",
        "Semua skill dinilai level 0",
        "Semua level = 0 (Tidak Tahu)",
        "1. Isi form dengan semua level 0\n2. Submit\n3. Pilih target job lalu jalankan analisis",
        "Analisis tetap berjalan tanpa error; tidak ada skill terkirim sebagai skill dimiliki sehingga seluruh skill benchmark masuk gap (match_score = 0)",
    ),
    (
        "TC-ASS-03",
        "Skill dengan level >= 1 dikirim sebagai skill dimiliki",
        "Mis. JavaScript level 2 — Target: Software Development",
        "1. Isi beberapa skill dengan level 1, 2, 3\n2. Submit dan jalankan analisis\n3. Periksa daftar matched",
        "Skill dengan level >= 1 muncul pada daftar matched bila merupakan skill benchmark; skill level 0 tidak pernah muncul di matched",
    ),
    (
        "TC-ASS-04",
        "Kombinasi level campuran (1, 2, 3)",
        "Level campuran pada tiap subdomain",
        "1. Isi form dengan kombinasi level 1/2/3\n2. Submit dan jalankan analisis",
        "Proses berhasil tanpa error; gap_score + match_score = 100 dan hasil tampil lengkap",
    ),
    (
        "TC-ASS-05",
        "Submit ulang dengan penilaian berbeda",
        "Dua set penilaian berbeda",
        "1. Isi dan submit penilaian pertama, jalankan analisis\n2. Kembali ke form, ubah penilaian, submit lagi\n3. Jalankan analisis ulang",
        "Hasil diperbarui sesuai penilaian terbaru tanpa perlu me-reload halaman",
    ),
    (
        "TC-ASS-06",
        "Kesesuaian hasil assessment dengan manual check",
        "Skill & job yang sama: JavaScript, SQL — Software Development",
        "1. Di /manual-input tambahkan skill yang sama dengan level >= 1 di assessment, pilih job yang sama, jalankan analisis\n2. Bandingkan dengan hasil assessment",
        "gap_score, match_score, dan daftar gap/rekomendasi identik untuk input skill & job yang sama",
    ),
]

MAN_CASES = [
    (
        "TC-MAN-01",
        "Input skill dengan spasi di awal/akhir",
        "\"  Python  \"",
        "1. Ketik \"  Python  \" pada input manual\n2. Klik Tambah",
        "Skill di-trim: chip menampilkan \"Python\" tanpa spasi berlebih",
    ),
    (
        "TC-MAN-02",
        "Input huruf besar/kecil berbeda dianggap beda",
        "\"python\" lalu \"Python\"",
        "1. Tambahkan \"python\"\n2. Tambahkan \"Python\" lagi",
        "Kedua chip ditambahkan sebagai entri terpisah (dedup frontend case-sensitive); backend tetap memproses keduanya dengan normalisasi lowercase tanpa error",
    ),
    (
        "TC-MAN-03",
        "Input banyak skill (15-20 skill)",
        "15-20 skill berbeda",
        "1. Tambahkan 15-20 skill\n2. Pilih target job\n3. Jalankan analisis",
        "Semua chip tampil dan counter sesuai; hasil benchmark_count = 15 (nilai top_n default) tanpa error",
    ),
    (
        "TC-MAN-04",
        "Menghapus seluruh skill",
        "-",
        "1. Tambahkan beberapa skill\n2. Hapus semua chip satu per satu",
        "Setelah semua chip terhapus, bagian Tahap 2 (pemilih job) dan hasil analisis menghilang",
    ),
    (
        "TC-MAN-05",
        "Analisis ulang dengan target job berbeda",
        "Skill sama, job pertama lalu job kedua",
        "1. Jalankan analisis dengan job A\n2. Ganti target job ke B\n3. Jalankan analisis ulang",
        "Hasil diperbarui sesuai job B; tidak ada sisa data hasil job A pada tampilan",
    ),
]

SUMMARY_ROWS = [
    ("Pengujian API", len(API_CASES), "", "", ""),
    ("Pengujian Frontend (UI)", len(UI_CASES), "", "", ""),
    ("Kesesuaian Scan CV dengan Gap/Rekomendasi", len(INT_CASES), "", "", ""),
    ("Pengujian Self Assessment", len(ASS_CASES), "", "", ""),
    ("Pengujian Manual Check", len(MAN_CASES), "", "", ""),
    ("Total", len(API_CASES) + len(UI_CASES) + len(INT_CASES) + len(ASS_CASES) + len(MAN_CASES), "", "", ""),
    ("Persentase keberhasilan", "-", "", "", "%"),
]

ENV_ROWS = [
    ("Sistem Operasi", "Windows 11"),
    ("Backend API", "FastAPI 0.136.3 (uvicorn 0.49.0), Python 3.13.9, port 8000"),
    ("Frontend", "React 19 + Vite 8, port 5173 (proxy /api ke port 8000)"),
    ("Model ML", "scikit-learn 1.7.2, joblib (multi-label skill prediction)"),
    ("Parser CV", "PyMuPDF 1.28.0 + spaCy (en_core_web_sm)"),
    ("Alat uji API", "Swagger UI http://localhost:8000/docs (atau Postman/curl)"),
    ("Browser", "(diisi)"),
    ("Data uji", "PDF CV berisi skill, file .txt/.docx, PDF > 10 MB, PDF korup, PDF scan tanpa teks"),
]


def set_cell(cell, text, bold=False, size=9, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def add_test_table(doc, cases, col_widths=COL_WIDTHS):
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(COLUMNS):
        set_cell(table.rows[0].cells[idx], header, bold=True, size=9, fill=HEADER_FILL)
        table.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (cid, skenario, inp, langkah, hasil) in enumerate(cases):
        row = table.add_row()
        fill = ALT_FILL if i % 2 == 0 else None
        values = [cid, skenario, inp, langkah, hasil, "", "", ""]
        for idx, val in enumerate(values):
            set_cell(row.cells[idx], val, size=9, fill=fill)
    set_col_widths(table, col_widths)
    return table


def landscape_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(1.8)
    return section


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # ── Halaman judul ─────────────────────────────────────────────────────────
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PENGUJIAN BLACK BOX")
    run.bold = True
    run.font.size = Pt(24)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Aplikasi Skill Gap Detection — Deteksi Kesenjangan Kompetensi IT")
    run.font.size = Pt(14)

    doc.add_paragraph()
    for label in ["Nama Penguji   : .......................",
                  "NIM                : .......................",
                  "Program Studi  : .......................",
                  "Tanggal Uji      : ......................."]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.runs.clear()
        p.add_run(label).font.size = Pt(12)
    doc.add_page_break()

    # ── 1. Pendahuluan ────────────────────────────────────────────────────────
    doc.add_heading("1. Pendahuluan", level=1)
    doc.add_heading("1.1 Tujuan", level=2)
    doc.add_paragraph(
        "Pengujian black box dilakukan untuk memverifikasi fungsionalitas aplikasi Skill Gap Detection "
        "tanpa meninjau struktur internal kode (white box). Pengujian difokuskan pada perilaku sistem "
        "berdasarkan input yang diberikan dan output yang dihasilkan, baik pada sisi API maupun antarmuka "
        "pengguna (frontend)."
    )
    doc.add_heading("1.2 Objek Pengujian", level=2)
    doc.add_paragraph(
        "a. Backend API FastAPI (http://localhost:8000) dengan endpoint: /health, /job-categories, "
        "/detect-gap, /recommend, dan /parse-cv.\n"
        "b. Frontend React (http://localhost:5173) dengan halaman: Beranda, Scan CV, Input Manual, "
        "Self Assessment, dan Tentang.\n"
        "c. Kesesuaian hasil antar modul: kebenaran hubungan data antara hasil scan CV, deteksi gap, "
        "dan rekomendasi skill yang diperlukan, serta konsistensi antara self assessment dan manual check."
    )
    doc.add_heading("1.3 Metode", level=2)
    doc.add_paragraph(
        "Metode black box testing menggunakan teknik equivalence partitioning dan boundary value "
        "analysis, mencakup pengujian data valid, data tidak valid, dan nilai batas (boundary). Selain "
        "pengujian fungsional per endpoint dan per halaman, dilakukan pula pengujian kesesuaian "
        "(consistency testing) untuk memastikan hasil scan CV selaras dengan daftar skill yang "
        "diperlukan (gap) dan rekomendasi yang diberikan, serta konsistensi hasil antara modul "
        "self assessment dan manual check. Setiap kasus uji dicatat pada tabel dengan status PASS "
        "(sesuai hasil yang diharapkan) atau FAIL (tidak sesuai)."
    )

    # ── 2. Lingkungan Pengujian ───────────────────────────────────────────────
    doc.add_heading("2. Lingkungan Pengujian", level=1)
    env_table = doc.add_table(rows=1, cols=2)
    env_table.style = "Table Grid"
    set_cell(env_table.rows[0].cells[0], "Komponen", bold=True, size=10, fill=HEADER_FILL)
    set_cell(env_table.rows[0].cells[1], "Keterangan", bold=True, size=10, fill=HEADER_FILL)
    env_table.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    env_table.rows[0].cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (komponen, keterangan) in enumerate(ENV_ROWS):
        row = env_table.add_row()
        set_cell(row.cells[0], komponen, bold=True, size=10, fill=ALT_FILL if i % 2 == 0 else None)
        set_cell(row.cells[1], keterangan, size=10, fill=ALT_FILL if i % 2 == 0 else None)
    set_col_widths(env_table, [Cm(4.5), Cm(12.5)])

    doc.add_heading("2.1 Prosedur Menjalankan Aplikasi", level=2)
    for line in [
        "1. Jalankan backend: python run_api.py (server berjalan pada http://localhost:8000, "
        "dokumentasi API pada http://localhost:8000/docs).",
        "2. Jalankan frontend: cd frontend lalu npm run dev (aplikasi pada http://localhost:5173).",
        "3. Pastikan data model pada folder final_model/ sudah tersedia (dimuat otomatis saat server start).",
        "4. Siapkan file uji: PDF CV berisi skill, file non-PDF, PDF > 10 MB, PDF korup, dan PDF scan tanpa teks.",
        "5. Lakukan pengujian sesuai skenario pada tabel kasus uji, lalu catat hasil aktual dan status PASS/FAIL.",
    ]:
        doc.add_paragraph(line)

    # ── 3. Kasus uji API (landscape) ──────────────────────────────────────────
    landscape_section(doc)
    doc.add_heading("3. Kasus Uji API (Backend)", level=1)
    add_test_table(doc, API_CASES)

    # ── 4. Kasus uji Frontend ─────────────────────────────────────────────────
    doc.add_heading("4. Kasus Uji Frontend (UI)", level=1)
    add_test_table(doc, UI_CASES)

    # ── 5. Kasus uji kesesuaian Scan CV ↔ Gap/Rekomendasi ─────────────────────
    doc.add_heading("5. Kasus Uji Kesesuaian Scan CV dengan Gap/Rekomendasi", level=1)
    doc.add_paragraph(
        "Pengujian ini memverifikasi bahwa hasil ekstraksi skill dari CV (scan) selaras dengan daftar "
        "skill yang diperlukan (gap) dan rekomendasi belajar yang diberikan sistem, serta konsistensi "
        "perhitungan antar endpoint."
    )
    add_test_table(doc, INT_CASES)

    # ── 6. Kasus uji Self Assessment ───────────────────────────────────────────
    doc.add_heading("6. Kasus Uji Self Assessment", level=1)
    doc.add_paragraph(
        "Pengujian mendalam terhadap alur penilaian mandiri (self assessment), mencakup validitas "
        "form, pemetaan level penilaian ke input analisis, dan konsistensinya dengan modul lain."
    )
    add_test_table(doc, ASS_CASES)

    # ── 7. Kasus uji Manual Check ──────────────────────────────────────────────
    doc.add_heading("7. Kasus Uji Manual Check", level=1)
    doc.add_paragraph(
        "Pengujian mendalam terhadap alur input skill manual, mencakup penanganan input tepi "
        "(spasi, huruf besar/kecil, jumlah skill) dan perilaku antarmuka saat data berubah."
    )
    add_test_table(doc, MAN_CASES)

    # ── 8. Ringkasan Hasil ─────────────────────────────────────────────────────
    doc.add_heading("8. Ringkasan Hasil Pengujian", level=1)
    doc.add_paragraph(
        "Kolom Hasil Aktual, Status, dan Catatan diisi setelah pengujian dilakukan. "
        "Ringkasan di bawah dihitung berdasarkan hasil seluruh kasus uji."
    )
    summary = doc.add_table(rows=1, cols=5)
    summary.style = "Table Grid"
    headers = ["Kasus Uji", "Jumlah", "PASS", "FAIL", "Keterangan"]
    for idx, h in enumerate(headers):
        set_cell(summary.rows[0].cells[idx], h, bold=True, size=10, fill=HEADER_FILL)
        summary.rows[0].cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (nama, jumlah, pas, fail, ket) in enumerate(SUMMARY_ROWS):
        row = summary.add_row()
        fill = ALT_FILL if i % 2 == 0 else None
        set_cell(row.cells[0], nama, bold=True, size=10, fill=fill)
        set_cell(row.cells[1], str(jumlah), size=10, fill=fill)
        set_cell(row.cells[2], pas, size=10, fill=fill)
        set_cell(row.cells[3], fail, size=10, fill=fill)
        set_cell(row.cells[4], ket, size=10, fill=fill)
    set_col_widths(summary, [Cm(7.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(3.0)])

    doc.save(OUT_PATH)
    print(f"Dokumen berhasil dibuat: {OUT_PATH}")


if __name__ == "__main__":
    main()
